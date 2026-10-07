"""Render ResumeData to ATS-friendly DOCX and PDF, plus a PDF analysis report.

ATS rules followed:
- single column, no tables / text boxes / images / icons in the resume
- nothing important in headers/footers
- standard headings only, simple round bullets
- Calibri (DOCX) / Helvetica ≈ Arial (PDF), 10-11 pt, ~0.6" margins
- PDF built from structured text with reportlab → real selectable text layer
"""

from __future__ import annotations

import io
from typing import List, Optional, Tuple
from xml.sax.saxutils import escape

from .models import AnalysisReport, ResumeData


def _date_range(start: str, end: str) -> str:
    if start and end:
        return f"{start} - {end}"
    return start or end or ""


def _contact_line(resume: ResumeData) -> str:
    c = resume.contact
    return " | ".join(x for x in [c.email, c.phone, c.location, *c.links] if x)


def _section_order(fresher: bool) -> List[str]:
    if fresher:
        return ["Summary", "Education", "Skills", "Projects", "Experience", "Certifications"]
    return ["Summary", "Skills", "Experience", "Projects", "Education", "Certifications"]


def _skill_lines(resume: ResumeData) -> List[Tuple[str, str]]:
    """(label, comma-separated skills). Empty label = single plain line."""
    if resume.skill_groups:
        return [(g, ", ".join(v)) for g, v in resume.skill_groups.items() if v]
    return [("", ", ".join(resume.skills))] if resume.skills else []


# --------------------------------------------------------------------------- #
# DOCX                                                                        #
# --------------------------------------------------------------------------- #
def render_docx(resume: ResumeData, fresher: bool = False) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.05

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11.5)

    def bullet(text: str) -> None:
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    def entry_line(bold: str, rest: List[str]) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.add_run(bold).bold = True
        tail = [x for x in rest if x]
        if tail:
            p.add_run(" | " + " | ".join(tail))

    # Contact block (plain text in the body, not in the header)
    name_p = doc.add_paragraph()
    name_run = name_p.add_run(resume.contact.name or "Your Name")
    name_run.bold = True
    name_run.font.size = Pt(16)
    contact = _contact_line(resume)
    if contact:
        doc.add_paragraph(contact)

    for section_name in _section_order(fresher):
        if section_name == "Summary" and resume.summary:
            heading("Summary")
            doc.add_paragraph(resume.summary)
        elif section_name == "Skills" and _skill_lines(resume):
            heading("Skills")
            for label, items in _skill_lines(resume):
                p = doc.add_paragraph()
                if label:
                    p.add_run(f"{label}: ").bold = True
                p.add_run(items)
        elif section_name == "Experience" and resume.experience:
            heading("Experience")
            for e in resume.experience:
                entry_line(e.title or e.company, [e.company if e.title else "", e.location,
                                                  _date_range(e.start_date, e.end_date)])
                for b in e.bullets:
                    bullet(b)
        elif section_name == "Projects" and resume.projects:
            heading("Projects")
            for p in resume.projects:
                tech = f"Tech: {', '.join(p.tech_stack)}" if p.tech_stack else ""
                entry_line(p.name, [p.role, tech, p.link, _date_range(p.start_date, p.end_date)])
                for b in p.bullets:
                    bullet(b)
        elif section_name == "Education" and resume.education:
            heading("Education")
            for ed in resume.education:
                degree = f"{ed.degree} in {ed.field}" if ed.degree and ed.field else (ed.degree or ed.field)
                entry_line(degree or ed.institution,
                           [ed.institution if degree else "", _date_range(ed.start_date, ed.end_date), ed.grade])
        elif section_name == "Certifications" and resume.certifications:
            heading("Certifications")
            for c in resume.certifications:
                bullet(c)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF helpers                                                                 #
# --------------------------------------------------------------------------- #
_PDF_REPLACE = {"₹": "Rs. ", "→": "->", "←": "<-", "≥": ">=", "≤": "<=", "×": "x", "\u200b": ""}
_WINANSI_EXTRA = set("–—‘’“”•…€")


def _safe(text: str) -> str:
    """Make text safe for the built-in Helvetica font (WinAnsi) and reportlab markup."""
    for bad, good in _PDF_REPLACE.items():
        text = text.replace(bad, good)
    text = "".join(ch for ch in text if ord(ch) < 256 or ch in _WINANSI_EXTRA)
    return escape(text)


def _styles():
    from reportlab.lib.styles import ParagraphStyle

    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=12.6, spaceAfter=1)
    return {
        "base": base,
        "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=18),
        "contact": ParagraphStyle("contact", parent=base, fontSize=9.5, spaceAfter=4),
        "heading": ParagraphStyle("heading", parent=base, fontName="Helvetica-Bold", fontSize=11.5,
                                  leading=14, spaceBefore=8, spaceAfter=2),
        "entry": ParagraphStyle("entry", parent=base, spaceBefore=4),
        "bullet": ParagraphStyle("bullet", parent=base, leftIndent=0),
        "small": ParagraphStyle("small", parent=base, fontSize=8.5, leading=10.5),
    }


def _bullets(items: List[str], style):
    from reportlab.platypus import ListFlowable, ListItem, Paragraph

    return ListFlowable(
        [ListItem(Paragraph(_safe(t), style), leftIndent=12, value="•") for t in items],
        bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8,
    )


def render_pdf(resume: ResumeData, fresher: bool = False) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            title=f"{resume.contact.name} - Resume", author=resume.contact.name)
    flow = [Paragraph(_safe(resume.contact.name or "Your Name"), st["name"])]
    contact = _contact_line(resume)
    if contact:
        flow.append(Paragraph(_safe(contact), st["contact"]))

    def entry(bold: str, rest: List[str]) -> Paragraph:
        tail = " | ".join(_safe(x) for x in rest if x)
        return Paragraph(f"<b>{_safe(bold)}</b>" + (f" | {tail}" if tail else ""), st["entry"])

    for section_name in _section_order(fresher):
        if section_name == "Summary" and resume.summary:
            flow += [Paragraph("SUMMARY", st["heading"]), Paragraph(_safe(resume.summary), st["base"])]
        elif section_name == "Skills" and _skill_lines(resume):
            flow.append(Paragraph("SKILLS", st["heading"]))
            for label, items in _skill_lines(resume):
                prefix = f"<b>{_safe(label)}:</b> " if label else ""
                flow.append(Paragraph(prefix + _safe(items), st["base"]))
        elif section_name == "Experience" and resume.experience:
            flow.append(Paragraph("EXPERIENCE", st["heading"]))
            for e in resume.experience:
                flow.append(entry(e.title or e.company, [e.company if e.title else "", e.location,
                                                         _date_range(e.start_date, e.end_date)]))
                if e.bullets:
                    flow.append(_bullets(e.bullets, st["bullet"]))
        elif section_name == "Projects" and resume.projects:
            flow.append(Paragraph("PROJECTS", st["heading"]))
            for p in resume.projects:
                tech = f"Tech: {', '.join(p.tech_stack)}" if p.tech_stack else ""
                flow.append(entry(p.name, [p.role, tech, p.link, _date_range(p.start_date, p.end_date)]))
                if p.bullets:
                    flow.append(_bullets(p.bullets, st["bullet"]))
        elif section_name == "Education" and resume.education:
            flow.append(Paragraph("EDUCATION", st["heading"]))
            for ed in resume.education:
                degree = f"{ed.degree} in {ed.field}" if ed.degree and ed.field else (ed.degree or ed.field)
                flow.append(entry(degree or ed.institution,
                                  [ed.institution if degree else "", _date_range(ed.start_date, ed.end_date),
                                   ed.grade]))
        elif section_name == "Certifications" and resume.certifications:
            flow += [Paragraph("CERTIFICATIONS", st["heading"]), _bullets(resume.certifications, st["bullet"])]

    doc.build(flow)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Analysis report PDF                                                         #
# --------------------------------------------------------------------------- #
def render_report_pdf(before: AnalysisReport, after: Optional[AnalysisReport] = None,
                      job_title: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch, title="Resume Analysis Report")
    flow = [Paragraph("Resume Analysis Report", st["name"])]
    if job_title:
        flow.append(Paragraph(_safe(f"Target role: {job_title}"), st["base"]))
    score_line = f"Original score: <b>{before.overall_score:.1f}/100</b>"
    if after:
        score_line += (f" &nbsp; Tailored score: <b>{after.overall_score:.1f}/100</b> "
                       f"&nbsp; Change: <b>{after.overall_score - before.overall_score:+.1f}</b>")
    flow += [Paragraph(score_line, st["base"]),
             Paragraph(_safe(f"Confidence: {before.confidence}. {before.disclaimer}"), st["small"]),
             Spacer(1, 6)]

    def table(rows, widths):
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    flow.append(Paragraph("Category breakdown", st["heading"]))
    header = ["Category", "Max", "Original"] + (["Tailored"] if after else [])
    rows = [header]
    for i, c in enumerate(before.category_scores):
        row = [Paragraph(_safe(c.name), st["small"]), f"{c.weight:g}", f"{c.points:.1f}"]
        if after:
            row.append(f"{after.category_scores[i].points:.1f}")
        rows.append(row)
    flow.append(table(rows, [3.2 * inch, 0.7 * inch, 1 * inch] + ([1 * inch] if after else [])))

    flow.append(Paragraph("Top fixes (original resume)", st["heading"]))
    for f in before.top_fixes:
        flow.append(Paragraph(_safe(f"+{f.estimated_gain:.1f} pts — {f.title}: {f.detail}"), st["base"]))

    final = after or before
    flow.append(Paragraph("Requirement coverage" + (" (tailored)" if after else ""), st["heading"]))
    rows = [["Requirement", "Importance", "Status", "Evidence"]]
    for m in final.requirement_matches:
        rows.append([Paragraph(_safe(m.requirement), st["small"]), m.importance.value,
                     m.status.value.replace("_", " "), Paragraph(_safe(m.evidence[:160]), st["small"])])
    flow.append(table(rows, [1.6 * inch, 0.8 * inch, 1.3 * inch, 3.2 * inch]))

    flow.append(Paragraph("ATS audit", st["heading"]))
    for issue in final.ats_report.issues:
        flow.append(Paragraph(_safe(f"[{issue.severity.upper()}] {issue.message}"), st["small"]))
    if not final.ats_report.issues:
        flow.append(Paragraph("No format issues detected.", st["small"]))

    doc.build(flow)
    return buf.getvalue()