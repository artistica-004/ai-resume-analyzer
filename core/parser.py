"""Robust resume text extraction (PDF + DOCX) with parse-quality flags.

Why pdfplumber (not pypdf)?
- It exposes word coordinates, which lets us detect two-column layouts and
  read each column separately (pypdf interleaves columns line by line).
- It can detect tables and images, which feeds the ATS audit and the
  scanned-PDF check.
- Cost: ~3 MB extra (pdfminer.six). No system binaries needed.

parse_resume() NEVER returns None: it returns a ParsedResume or raises
ResumeParseError with a message that is safe to show to the user.
"""

from __future__ import annotations

import io
import logging
import math
import re
from pathlib import Path
from typing import Any, List, Optional, Tuple

from .models import ParsedResume, ParseQuality

logger = logging.getLogger(__name__)

MAX_FILE_MB = 5
MIN_TEXT_CHARS = 200
MAX_PDF_PAGES_READ = 10
ALLOWED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx"}

_BULLET_CHARS = "•●▪■◦○◆◇►▸➢➤✓✔✦·*\uf0b7\uf0a7\uf0d8\uf076\u2022\u2023\u2043"
_LIGATURES = {"\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi", "\ufb04": "ffl"}
_CONTACT_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+|\+?\d[\d\s().-]{8,}\d")


class ResumeParseError(Exception):
    """User-facing parsing error (message is safe to display)."""


# --------------------------------------------------------------------------- #
# Validation                                                                  #
# --------------------------------------------------------------------------- #
def validate_upload(file_name: str, data: bytes) -> str:
    """Check extension, size and magic bytes. Returns 'pdf' or 'docx'."""
    ext = Path(file_name).suffix.lower()
    if ext == ".doc":
        raise ResumeParseError(
            "Old .doc files are not supported. Open it in Word and 'Save As' .docx, or export to PDF."
        )
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX resume.")
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    size_mb = len(data) / (1024 * 1024)
    if size_mb > MAX_FILE_MB:
        raise ResumeParseError(f"File is {size_mb:.1f} MB. The maximum allowed size is {MAX_FILE_MB} MB.")
    file_type = ALLOWED_EXTENSIONS[ext]
    if file_type == "pdf" and not data[:1024].lstrip().startswith(b"%PDF"):
        raise ResumeParseError("This file has a .pdf extension but is not a valid PDF.")
    if file_type == "docx" and not data.startswith(b"PK"):
        raise ResumeParseError("This file has a .docx extension but is not a valid Word document.")
    return file_type


# --------------------------------------------------------------------------- #
# Text cleaning                                                               #
# --------------------------------------------------------------------------- #
def clean_text(text: str) -> str:
    """Fix ligatures, unify bullet characters to '- ', collapse whitespace."""
    for bad, good in _LIGATURES.items():
        text = text.replace(bad, good)
    text = re.sub(r"\(cid:\d+\)", " ", text)
    text = text.replace("\u00a0", " ").replace("\t", " ")

    lines: List[str] = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        if s[0] in _BULLET_CHARS or re.match(r"^[-–—]\s+", s):
            s = "- " + s.lstrip(_BULLET_CHARS + "-–— ").strip()
        s = re.sub(r" {2,}", " ", s)
        lines.append(s)
    return "\n".join(lines).strip()


# --------------------------------------------------------------------------- #
# PDF                                                                         #
# --------------------------------------------------------------------------- #
def _find_column_split(page: Any) -> Optional[float]:
    """Return an x-coordinate that splits a two-column page, or None.

    Heuristic: find a vertical line in the middle 30-70% of the page that
    almost no word crosses, with plenty of words on both sides, and where
    few text lines 'bridge' it with a normal word gap (which would indicate
    ordinary single-column text).
    """
    words = page.extract_words()
    if len(words) < 60:
        return None
    x0, _, x1, _ = page.bbox
    width = x1 - x0

    best_x: Optional[float] = None
    best_cross: Optional[int] = None
    for pct in range(30, 71, 2):
        x = x0 + width * pct / 100
        crossing = sum(1 for w in words if w["x0"] < x < w["x1"])
        left = sum(1 for w in words if w["x1"] <= x)
        right = len(words) - left - crossing
        if left < 20 or right < 20:
            continue
        if best_cross is None or crossing < best_cross:
            best_x, best_cross = x, crossing

    if best_x is None or best_cross is None or best_cross > max(2, int(len(words) * 0.01)):
        return None

    # Group words into visual lines and count lines that bridge the split
    lines: dict = {}
    for w in words:
        lines.setdefault(round(w["top"] / 3), []).append(w)
    bridged = 0
    for line_words in lines.values():
        lefts = [w for w in line_words if w["x1"] <= best_x]
        rights = [w for w in line_words if w["x0"] >= best_x]
        if lefts and rights:
            gap = min(w["x0"] for w in rights) - max(w["x1"] for w in lefts)
            if gap < 12:  # normal space between words -> same column
                bridged += 1
    if bridged > len(lines) * 0.15:
        return None
    return best_x


def _parse_pdf(data: bytes) -> Tuple[str, ParseQuality]:
    import pdfplumber

    quality = ParseQuality(method="pdfplumber")
    page_texts: List[str] = []
    scanned_pages = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            quality.num_pages = len(pdf.pages)
            if quality.num_pages == 0:
                raise ResumeParseError("The PDF has no pages.")
            if quality.num_pages > MAX_PDF_PAGES_READ:
                quality.warnings.append(
                    f"Only the first {MAX_PDF_PAGES_READ} of {quality.num_pages} pages were read."
                )
            for page in pdf.pages[:MAX_PDF_PAGES_READ]:
                if page.images:
                    quality.has_images = True
                try:
                    if page.find_tables():
                        quality.has_tables = True
                except Exception:  # noqa: BLE001 - table detection is best-effort
                    pass

                split = _find_column_split(page)
                if split is not None:
                    quality.multi_column_suspected = True
                    bx0, top, bx1, bottom = page.bbox
                    left = page.crop((bx0, top, split, bottom)).extract_text() or ""
                    right = page.crop((split, top, bx1, bottom)).extract_text() or ""
                    text = f"{left}\n{right}"
                else:
                    text = page.extract_text() or ""

                if len(text.strip()) < 20 and page.images:
                    scanned_pages += 1
                page_texts.append(text)
    except ResumeParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF could not be opened: %s", type(exc).__name__)
        raise ResumeParseError(
            "This PDF could not be read. It may be corrupted or password-protected. "
            "Try re-exporting it, or upload a DOCX instead."
        ) from exc

    pages_read = min(quality.num_pages, MAX_PDF_PAGES_READ)
    if scanned_pages and scanned_pages == pages_read:
        quality.is_scanned = True
    elif scanned_pages:
        quality.warnings.append(
            f"{scanned_pages} page(s) look like images with no text layer; their content was skipped."
        )
    return "\n\n".join(page_texts), quality


# --------------------------------------------------------------------------- #
# DOCX                                                                        #
# --------------------------------------------------------------------------- #
def _docx_paragraph_text(paragraph: Any) -> str:
    """Paragraph text, prefixed with '- ' if it is a list/bullet item."""
    text = paragraph.text.strip()
    if not text:
        return ""
    is_list = False
    try:
        style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
        is_list = "list" in style_name.lower()
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.numPr is not None:
            is_list = True
    except Exception:  # noqa: BLE001
        pass
    return f"- {text}" if is_list and not text.startswith("-") else text


def _docx_table_lines(table: Any) -> List[str]:
    """Read a table row by row, skipping repeated merged cells."""
    lines: List[str] = []
    for row in table.rows:
        seen = set()
        for cell in row.cells:
            cell_id = id(cell._tc)
            if cell_id in seen:
                continue
            seen.add(cell_id)
            for para in cell.paragraphs:
                t = _docx_paragraph_text(para)
                if t:
                    lines.append(t)
    return lines


def _parse_docx(data: bytes) -> Tuple[str, ParseQuality]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    quality = ParseQuality(method="python-docx")
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        logger.warning("DOCX could not be opened: %s", type(exc).__name__)
        raise ResumeParseError(
            "This Word document could not be read. Try re-saving it as .docx or exporting to PDF."
        ) from exc

    body_lines: List[str] = []
    # Walk the body in document order so tables stay where they appear
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = _docx_paragraph_text(Paragraph(child, doc))
            if t:
                body_lines.append(t)
        elif child.tag == qn("w:tbl"):
            quality.has_tables = True
            body_lines.extend(_docx_table_lines(Table(child, doc)))

    # Text boxes (python-docx's paragraph.text does not include them)
    textbox_lines: List[str] = []
    seen_tb = set()
    for tb in doc.element.body.iter(qn("w:txbxContent")):
        for p in tb.iter(qn("w:p")):
            t = "".join(node.text or "" for node in p.iter(qn("w:t"))).strip()
            if t and t not in seen_tb:
                seen_tb.add(t)
                textbox_lines.append(t)
    if textbox_lines:
        quality.warnings.append(
            "Text boxes were found. Many ATS systems skip text-box content; move it into normal paragraphs."
        )

    # Headers and footers
    hf_lines: List[str] = []
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                if part.is_linked_to_previous:
                    continue
                for para in part.paragraphs:
                    t = para.text.strip()
                    if t and t not in hf_lines:
                        hf_lines.append(t)
                for table in part.tables:
                    for t in _docx_table_lines(table):
                        if t not in hf_lines:
                            hf_lines.append(t)
            except Exception:  # noqa: BLE001
                continue
    if hf_lines:
        if any(_CONTACT_RE.search(line) for line in hf_lines):
            quality.header_footer_text = True
            quality.warnings.append(
                "Contact details are in the page header/footer. Some ATS systems ignore headers; "
                "put them in the main body."
            )

    if doc.inline_shapes and len(doc.inline_shapes) > 0:
        quality.has_images = True

    all_lines = hf_lines + body_lines + textbox_lines
    text = "\n".join(all_lines)
    words = len(text.split())
    quality.num_pages = max(1, math.ceil(words / 600))  # estimate: DOCX has no fixed pages
    return text, quality


# --------------------------------------------------------------------------- #
# Public API                                                                  #
# --------------------------------------------------------------------------- #
def _finalize(raw_text: str, file_name: str, file_type: str, quality: ParseQuality) -> ParsedResume:
    text = clean_text(raw_text)
    quality.char_count = len(text)
    quality.word_count = len(text.split())

    if quality.is_scanned or (file_type == "pdf" and quality.char_count < MIN_TEXT_CHARS and quality.has_images):
        quality.is_scanned = True
        raise ResumeParseError(
            "This looks like a scanned or image-only PDF, so no text could be extracted. "
            "Please upload the original DOCX, or export a text-based PDF from Word/Google Docs. "
            "(If you only have a scan, run it through an OCR tool first, e.g. Google Drive "
            "'Open with Google Docs'.)"
        )
    if quality.char_count < MIN_TEXT_CHARS:
        raise ResumeParseError(
            f"Only {quality.char_count} characters of text were found. The file may be empty, "
            "image-based, or heavily formatted. Please try a simpler PDF or a DOCX."
        )

    if quality.multi_column_suspected:
        quality.warnings.append(
            "A multi-column layout was detected. We read each column separately, but many ATS "
            "systems scramble columns; a single-column layout is safer."
        )
    if quality.has_tables:
        quality.warnings.append("Tables were detected. Some ATS systems read tables in the wrong order.")

    letters = [c for c in text if c.isalpha()]
    if letters and sum(c.isascii() for c in letters) / len(letters) < 0.7:
        quality.warnings.append("The resume does not appear to be in English; results may be less accurate.")

    logger.info("Parsed %s resume: %d chars, %d pages", file_type, quality.char_count, quality.num_pages)
    return ParsedResume(raw_text=text, file_name=file_name, file_type=file_type, quality=quality)


def parse_resume_bytes(data: bytes, file_name: str) -> ParsedResume:
    """Parse raw bytes of a PDF/DOCX file."""
    file_type = validate_upload(file_name, data)
    if file_type == "pdf":
        raw, quality = _parse_pdf(data)
    else:
        raw, quality = _parse_docx(data)
    return _finalize(raw, file_name, file_type, quality)


def parse_resume(uploaded_file: Any) -> ParsedResume:
    """Parse a Streamlit UploadedFile, a file path (str/Path), or any object
    with .name and .getvalue()/.read().
    """
    if uploaded_file is None:
        raise ResumeParseError("Please upload a resume file first.")
    if isinstance(uploaded_file, (str, Path)):
        path = Path(uploaded_file)
        if not path.exists():
            raise ResumeParseError(f"File not found: {path.name}")
        return parse_resume_bytes(path.read_bytes(), path.name)

    name = getattr(uploaded_file, "name", "resume")
    if hasattr(uploaded_file, "getvalue"):
        data = uploaded_file.getvalue()
    else:
        data = uploaded_file.read()
    return parse_resume_bytes(data, name)