"""Create the two sample resumes (deliberately imperfect, to demo improvement).

Run: python scripts/make_samples.py
- sample_data/resume_fresher.docx    : odd heading, skills in a table, weak bullets
- sample_data/resume_experienced.pdf : 4-year backend dev, vague bullets, mixed dates
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parents[1] / "sample_data"


def make_fresher_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Ravi Kumar")
    doc.add_paragraph("ravi.kumar@example.com | +91 9123456780 | Vijayawada | github.com/ravikumar")
    doc.add_paragraph("CAREER OBJECTIVE")
    doc.add_paragraph("To obtain a challenging position in a reputed organization where I can use my skills.")
    doc.add_paragraph("TECHNICAL SKILLS")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "Python, SQL, JS"
    table.rows[1].cells[0].text = "Libraries"
    table.rows[1].cells[1].text = "Pandas, NumPy, scikit-learn, Flask, Git"
    doc.add_paragraph("MY JOURNEY")
    doc.add_paragraph("Data Science Intern, InsightWorks Analytics (May 2024 - Jul 2024)")
    for b in ["Responsible for cleaning sales data using pandas",
              "Worked on a model to predict customer churn with scikit-learn that got 84% accuracy",
              "Helped the team in making reports"]:
        doc.add_paragraph(b, style="List Bullet")
    doc.add_paragraph("PROJECTS")
    doc.add_paragraph("Movie Recommender (Python, scikit-learn, Flask)")
    for b in ["Made a movie recommendation system using collaborative filtering on 100k ratings",
              "Created rest api using flask to give recommendations"]:
        doc.add_paragraph(b, style="List Bullet")
    doc.add_paragraph("Fake News Detector (Python, NLP, scikit-learn)")
    doc.add_paragraph("Built a text classifier with TF-IDF and logistic regression", style="List Bullet")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("B.Tech in Computer Science, Sri Venkateswara Engineering College, 2021 - 2025, CGPA 8.2/10")
    doc.add_paragraph("CERTIFICATIONS")
    doc.add_paragraph("Machine Learning Specialization (Coursera)", style="List Bullet")
    doc.save(path)


def make_experienced_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    body, h = styles["BodyText"], styles["Heading3"]
    story = [
        Paragraph("<b>Neha Sharma</b>", styles["Title"]),
        Paragraph("neha.sharma@example.com | +91 9988776655 | Pune | linkedin.com/in/nehasharma", body),
        Paragraph("PROFILE", h),
        Paragraph("Software developer with experience in Python web development.", body),
        Paragraph("SKILLS", h),
        Paragraph("Python, Django, Flask, REST APIs, PostgreSQL, MySQL, Redis, Docker, Git, Jenkins, Linux", body),
        Paragraph("WORK EXPERIENCE", h),
        Paragraph("<b>Software Engineer</b>, PayBridge Technologies, 06/2022 - Present", body),
        Paragraph("• Was responsible for developing APIs for the payments module using Django", body),
        Paragraph("• Worked on PostgreSQL queries and improved performance", body),
        Paragraph("• Used Redis caching to reduce API latency by 40%", body),
        Paragraph("• Wrote unit tests using pytest", body),
        Paragraph("<b>Associate Developer</b>, Codecraft Solutions, Jan 2021 - May 2022", body),
        Paragraph("• Developed internal tools in Flask for 200+ employees", body),
        Paragraph("• Helped in moving the deployment to Docker containers", body),
        Paragraph("• Fixed bugs", body),
        Paragraph("PROJECTS", h),
        Paragraph("<b>Expense Tracker API</b> (FastAPI, PostgreSQL)", body),
        Paragraph("• Built an expense tracking REST API with JWT authentication", body),
        Paragraph("EDUCATION", h),
        Paragraph("B.E. Computer Engineering, Pune Institute of Technology, 2016 - 2020", body),
        Spacer(1, 6),
    ]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_fresher_docx(OUT / "resume_fresher.docx")
    make_experienced_pdf(OUT / "resume_experienced.pdf")
    print("Created:", OUT / "resume_fresher.docx", "and", OUT / "resume_experienced.pdf")