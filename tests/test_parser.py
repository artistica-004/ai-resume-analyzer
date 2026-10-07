"""Parser: DOCX/PDF extraction, tables, headers, scanned detection, validation."""

import io
from pathlib import Path

import pytest
from docx import Document

from core.parser import ResumeParseError, clean_text, parse_resume, parse_resume_bytes, validate_upload
from core.renderer import render_docx, render_pdf

SAMPLES = Path(__file__).resolve().parents[1] / "sample_data"
LONG = ("Built data pipelines and REST APIs in Python for analytics dashboards used by internal teams. " * 4)


def test_rendered_docx_roundtrip(resume):
    parsed = parse_resume_bytes(render_docx(resume), "r.docx")
    assert parsed.file_type == "docx"
    assert "asha@example.com" in parsed.raw_text
    assert "DataNest Labs" in parsed.raw_text
    assert "- Built a churn prediction model" in parsed.raw_text  # bullets normalised
    assert not parsed.quality.has_tables


def test_rendered_pdf_has_text_layer(resume):
    parsed = parse_resume_bytes(render_pdf(resume), "r.pdf")
    assert parsed.file_type == "pdf"
    assert "DataNest Labs" in parsed.raw_text
    assert parsed.quality.num_pages == 1
    assert not parsed.quality.is_scanned


def test_docx_tables_are_read():
    doc = Document()
    doc.add_paragraph("Asha Rao | asha@example.com")
    doc.add_paragraph(LONG)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "Docker"
    buf = io.BytesIO()
    doc.save(buf)
    parsed = parse_resume_bytes(buf.getvalue(), "t.docx")
    assert parsed.quality.has_tables
    assert "Docker" in parsed.raw_text


def test_docx_header_contact_is_read_and_flagged():
    doc = Document()
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "asha@example.com | +91 9876543210"
    doc.add_paragraph(LONG)
    buf = io.BytesIO()
    doc.save(buf)
    parsed = parse_resume_bytes(buf.getvalue(), "h.docx")
    assert parsed.quality.header_footer_text
    assert "asha@example.com" in parsed.raw_text


def test_scanned_pdf_is_detected():
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (600, 800), "white")
    ImageDraw.Draw(img).rectangle([50, 50, 550, 120], fill="black")
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawImage(ImageReader(img), 30, 30, width=500, height=700)
    c.save()
    with pytest.raises(ResumeParseError, match="scanned"):
        parse_resume_bytes(buf.getvalue(), "scan.pdf")


@pytest.mark.parametrize("name,data", [
    ("resume.doc", b"anything"),
    ("resume.txt", b"plain text"),
    ("resume.pdf", b"not really a pdf"),
    ("resume.docx", b"not a zip"),
    ("resume.pdf", b""),
    ("resume.pdf", b"%PDF" + b"0" * (6 * 1024 * 1024)),
], ids=["legacy-doc", "txt", "fake-pdf", "fake-docx", "empty", "oversize-6mb"])
def test_invalid_uploads_rejected(name, data):
    with pytest.raises(ResumeParseError):
        validate_upload(name, data)


def test_none_upload_raises_clear_error():
    with pytest.raises(ResumeParseError):
        parse_resume(None)


def test_clean_text_normalises_bullets():
    assert clean_text("• Built X\n▪ Did Y\n\n\n– Led Z") == "- Built X\n- Did Y\n\n- Led Z"


@pytest.mark.parametrize("name", ["resume_fresher.docx", "resume_experienced.pdf"])
def test_sample_files_parse(name):
    path = SAMPLES / name
    if not path.exists():
        pytest.skip("Run `python scripts/make_samples.py` first")
    parsed = parse_resume(path)
    assert parsed.quality.char_count > 500